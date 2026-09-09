"""Shared vocabulary of the ODCS Excel workbook: element references and cell value typing.

An element reference names the contract element a row on a child sheet belongs to: an
`Scope` and a `Scope Name`, resolved through the element's natural key (see
datacontract.model.natural_keys) or its `id` where it has no natural key. There is no positional
fallback: an element whose key is missing or duplicated cannot be referenced.
"""

import io
import json
import logging
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from typing import Any, Iterator, Optional

import openpyxl
import yaml
from open_data_contract_standard.model import OpenDataContractStandard, SchemaProperty, Team

from datacontract.model.natural_keys import NATURAL_KEYS

logger = logging.getLogger(__name__)

# the merged header above the inline custom property columns; each column is named after one property.
# Matched as a prefix: the row sheets label it "Custom Properties", Servers "Custom Properties (add as needed)".
CUSTOM_PROPERTIES_GROUP = "Custom Properties"

SERVER_FIELDS = [
    "account",
    "catalog",
    "catalogUrl",
    "database",
    "dataset",
    "delimiter",
    "encoding",
    "endpointUrl",
    "format",
    "host",
    "location",
    "namespace",
    "path",
    "port",
    "project",
    "region",
    "regionName",
    "schema",
    "serviceName",
    "stagingDir",
    "stream",
    "warehouse",
    "workgroup",
]


def server_field_name(workbook, server_type: Optional[str], field: str) -> Optional[str]:
    """The named range holding a server field: the generic block, else the legacy per-type block, else the legacy custom block."""
    legacy_type = "postgres" if server_type == "postgresql" else server_type
    for name in (f"servers.{field}", f"servers.{legacy_type}.{field.lower()}", f"servers.custom.{field}"):
        if name in workbook.defined_names:
            return name
    return None


@dataclass
class Element:
    kind: str
    ref: Optional[str]  # None when the element cannot be referenced
    obj: Any
    label: str  # human-readable, for warnings
    unaddressable_reason: Optional[str] = None


def iter_elements(odcs: OpenDataContractStandard) -> Iterator[Element]:
    """Every contract element that can carry custom properties or authoritative definitions."""
    yield Element("Contract", "", odcs, "the contract")
    if odcs.description:
        yield Element("Description", "", odcs.description, "the description")
    for server in _keyed(odcs.servers, NATURAL_KEYS["servers"], "Server", "server"):
        yield server
    for schema in _keyed(odcs.schema_, NATURAL_KEYS["schema"], "Schema", "schema"):
        yield schema
        if schema.obj.context is not None and not isinstance(schema.obj.context, str):
            yield from _context_elements(schema.obj.context, schema.label)
        for synonym in _synonyms(schema.obj, schema.ref, schema.label):
            yield synonym
        for relationship in _by_id(schema.obj.relationships, "Relationship", f"relationship of {schema.label}"):
            yield relationship
        for quality in _by_id(schema.obj.quality, "Quality", f"quality rule of {schema.label}"):
            yield quality
        yield from _property_elements(schema.obj.properties, schema.ref, schema.label)
    for support in _keyed(odcs.support, NATURAL_KEYS["support"], "Support", "support channel"):
        yield support
    if isinstance(odcs.team, Team):
        yield Element("Team", "", odcs.team, "the team")
        members = odcs.team.members
    else:
        members = odcs.team
    for member in _keyed(members, NATURAL_KEYS["team.members"], "Team Member", "team member"):
        yield member
    for role in _keyed(odcs.roles, NATURAL_KEYS["roles"], "Role", "role"):
        yield role
    for sla in _keyed(odcs.slaProperties, NATURAL_KEYS["slaProperties"], "SLA", "SLA property"):
        yield sla
    if odcs.context is not None and not isinstance(odcs.context, str):
        yield from _context_elements(odcs.context, "the contract")


def _property_elements(properties, schema_ref, schema_label, prefix="") -> Iterator[Element]:
    for prop in _keyed(properties, NATURAL_KEYS["schema.properties"], "Property", f"property of {schema_label}"):
        path = None if prop.ref is None else f"{prefix}{prop.ref}"
        prop.ref = None if path is None or schema_ref is None else f"{schema_ref}.{path}"
        prop.label = f"property {path or '?'} of {schema_label}"
        yield from _property_element_and_children(prop, schema_ref, schema_label, path)


def _property_element_and_children(prop: Element, schema_ref, schema_label, path) -> Iterator[Element]:
    yield prop
    p: SchemaProperty = prop.obj
    for enum_value in _keyed(p.enum, NATURAL_KEYS["enum"], "Enum Value", f"enum value of {prop.label}"):
        enum_value.ref = None if prop.ref is None or enum_value.ref is None else f"{prop.ref}={enum_value.ref}"
        yield enum_value
    yield from _synonyms(p, prop.ref, prop.label)
    yield from _by_id(p.relationships, "Relationship", f"relationship of {prop.label}")
    yield from _by_id(p.quality, "Quality", f"quality rule of {prop.label}")
    if p.properties:
        yield from _property_elements(p.properties, schema_ref, schema_label, prefix=f"{path or '?'}.")
    if p.items:
        # array items are written as "<parent>.items" in the schema sheet, no name needed
        items_path = None if path is None else f"{path}.items"
        items = Element(
            "Property", None if prop.ref is None else f"{prop.ref}.items", p.items, f"items of {prop.label}"
        )
        yield from _property_element_and_children(items, schema_ref, schema_label, items_path)


def _synonyms(owner, owner_ref, owner_label) -> Iterator[Element]:
    for synonym in _keyed(owner.synonyms, NATURAL_KEYS["synonyms"], "Synonym", f"synonym of {owner_label}"):
        synonym.ref = None if owner_ref is None or synonym.ref is None else f"{owner_ref}={synonym.ref}"
        yield synonym


def _context_elements(context, owner_label) -> Iterator[Element]:
    yield from _by_id(context.verifiedStatements, "Verified Statement", f"verified statement of {owner_label}")
    yield from _by_id(context.constraints, "Constraint", f"constraint of {owner_label}")


def _keyed(items, key_field, kind, noun) -> list[Element]:
    """Elements of a list keyed by a natural key; missing or duplicated keys make an element unreferenceable."""
    if not items:
        return []
    keys = [getattr(item, key_field, None) for item in items]
    keys = [None if k is None or str(k) == "" else str(k) for k in keys]
    elements = []
    for item, key in zip(items, keys):
        if key is None:
            elements.append(Element(kind, None, item, f"{noun} without {key_field}", f"has no {key_field}"))
        elif keys.count(key) > 1:
            elements.append(Element(kind, None, item, f"{noun} {key}", f"shares its {key_field} with another"))
        else:
            elements.append(Element(kind, key, item, f"{noun} {key}"))
    return elements


def _by_id(items, kind, noun) -> list[Element]:
    if not items:
        return []
    ids = [getattr(item, "id", None) or None for item in items]
    elements = []
    for item, item_id in zip(items, ids):
        if item_id is None:
            elements.append(Element(kind, None, item, noun, "has no id"))
        elif ids.count(item_id) > 1:
            elements.append(Element(kind, None, item, f"{noun} {item_id}", "shares its id with another"))
        else:
            elements.append(Element(kind, item_id, item, f"{noun} {item_id}"))
    return elements


def element_index(odcs: OpenDataContractStandard) -> dict[tuple[str, str], Element]:
    return {(e.kind, e.ref): e for e in iter_elements(odcs) if e.ref is not None}


def resolve_cell_value(value: Any, type_hint: Optional[str] = None) -> Any:
    """The value a workbook cell stands for.

    A blank `Type` resolves a text cell as YAML would (`true` is a boolean, `42` an integer); typed
    cells keep their type. `Text` is verbatim; `JSON` parses arrays and objects.
    """
    if value is None:
        return None
    hint = (type_hint or "").strip().lower()
    if hint == "text":
        return str(value)
    if hint == "json":
        return json.loads(str(value))
    if isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    text = str(value).strip()
    if text == "":
        return None
    try:
        resolved = yaml.safe_load(text)
    except yaml.YAMLError:
        return text
    if resolved is None or isinstance(resolved, (bool, int, float, str)):
        return resolved
    return text


def is_scalar(value: Any) -> bool:
    return value is None or isinstance(value, (bool, int, float, str))


class InlineRoundTrip:
    """Decides at export time whether a scalar survives a plain cell: written, saved, read back, resolved."""

    def __init__(self, values):
        self._ok: dict[tuple[str, Any], bool] = {}
        candidates = list({_key(v): v for v in values if is_scalar(v)}.values())
        if not candidates:
            return
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        for row, value in enumerate(candidates, start=1):
            sheet.cell(row=row, column=1, value=value)
        buffer = io.BytesIO()
        workbook.save(buffer)
        read_back = openpyxl.load_workbook(io.BytesIO(buffer.getvalue()), data_only=True).active
        for row, value in enumerate(candidates, start=1):
            resolved = resolve_cell_value(read_back.cell(row=row, column=1).value)
            self._ok[_key(value)] = resolved == value and type(resolved) is type(value)

    def survives(self, value: Any) -> bool:
        return is_scalar(value) and self._ok.get(_key(value), False)


def _key(value):
    return (type(value).__name__, value)
