"""Writes ../delivery.xlsx, kept as a script so the workbook can be reviewed and rebuilt."""

import datetime as dt
from pathlib import Path

from openpyxl import Workbook

ORDERS = [
    ["order_id", "customer_id", "order_date", "order_timestamp", "amount", "quantity", "paid", "note"],
    ["A1", "00123", dt.date(2026, 1, 2), dt.datetime(2026, 1, 2, 13, 45), 10.5, 2, True, "first"],
    ["A2", "00456", dt.date(2026, 1, 3), dt.datetime(2026, 1, 3, 8, 0), 7, 1, False, None],
    ["A3", "00789", dt.date(2026, 1, 5), dt.datetime(2026, 1, 5, 0, 0), 3.25, 5, True, "third"],
]
COUNTRY_CODES = [["code", "name"], ["DE", "Germany"], ["FR", "France"]]

workbook = Workbook()
workbook.remove(workbook.active)
for title, rows in {"Orders": ORDERS, "Country Codes": COUNTRY_CODES}.items():
    sheet = workbook.create_sheet(title)
    for row in rows:
        sheet.append(row)
    # Excel stores dates as numbers; only the format makes them dates
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            if isinstance(cell.value, dt.datetime):
                cell.number_format = "yyyy-mm-dd hh:mm"
            elif isinstance(cell.value, dt.date):
                cell.number_format = "yyyy-mm-dd"
workbook.save(Path(__file__).parent.parent / "delivery.xlsx")
