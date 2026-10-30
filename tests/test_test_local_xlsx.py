"""`datacontract test` against Excel workbooks, read through DuckDB's excel extension."""

import datetime as dt
import shutil
from pathlib import Path

import pytest
import yaml
from openpyxl import load_workbook

from datacontract.data_contract import DataContract
from datacontract.engines.ibis.connections.duckdb_connection import get_duckdb_connection
from datacontract.lint import resolve
from datacontract.model.run import Run

CONTRACT = "fixtures/xlsx/datacontract.yaml"
WORKBOOK = "fixtures/xlsx/delivery.xlsx"


def contract(path: str, models: list[str] | None = None, **order_changes) -> str:
    data = yaml.safe_load(Path(CONTRACT).read_text())
    data["servers"][0]["path"] = path
    if models is not None:
        data["schema"] = [s for s in data["schema"] if s["name"] in models]
    data["schema"][0].update(order_changes)
    return yaml.safe_dump(data)


def edited(tmp_path: Path, sheet_title: str = "Orders", **cells) -> str:
    """The fixture workbook with cells of its Orders sheet changed, e.g. E3="n/a"."""
    workbook = load_workbook(WORKBOOK)
    orders = workbook["Orders"]
    for cell, value in cells.items():
        orders[cell] = value
        if isinstance(value, dt.date):
            orders[cell].number_format = "yyyy-mm-dd"
    orders.title = sheet_title
    workbook.save(tmp_path / "delivery.xlsx")
    return str(tmp_path / "delivery.xlsx")


def failed(run) -> dict:
    return {check.name: check.reason for check in run.checks if check.result is None or check.result.value != "passed"}


def rows(contract_str: str, view: str) -> list[tuple]:
    data_contract = resolve.resolve_data_contract(data_contract_str=contract_str)
    con = get_duckdb_connection(data_contract, data_contract.servers[0], Run.create_run())
    return con.sql(f'SELECT * FROM "{view}"').fetchall()


def test_both_sheets_of_a_workbook_are_checked():
    run = DataContract(data_contract_file=CONTRACT).test()

    assert run.result == "passed", [(c.name, c.reason) for c in run.checks if c.result != "passed"]
    assert {"Orders", "Country Codes"} <= {check.model for check in run.checks}


def test_cells_are_read_with_the_types_of_the_contract():
    first, second, _ = rows(contract(WORKBOOK), "orders")

    assert first == ("A1", "00123", dt.date(2026, 1, 2), dt.datetime(2026, 1, 2, 13, 45), 10.5, 2, True, "first")
    assert second[-1] is None


def test_a_model_without_physical_name_reads_the_first_sheet():
    data = yaml.safe_load(contract(WORKBOOK, models=["orders"]))
    del data["schema"][0]["physicalName"]

    run = DataContract(data_contract_str=yaml.safe_dump(data)).test()

    assert run.result == "passed", failed(run)


def test_models_that_share_a_workbook_must_name_their_sheet():
    data = yaml.safe_load(contract(WORKBOOK))
    del data["schema"][1]["physicalName"]

    run = DataContract(data_contract_str=yaml.safe_dump(data)).test()

    assert run.result != "passed"
    assert any("physicalName" in (c.reason or "") and "country_codes" in (c.reason or "") for c in run.checks)


def test_a_glob_path_is_refused(tmp_path: Path):
    for name in ("a.xlsx", "b.xlsx"):
        shutil.copy(WORKBOOK, tmp_path / name)

    run = DataContract(data_contract_str=contract(str(tmp_path / "*.xlsx"), models=["orders"])).test()

    assert run.result != "passed"
    assert any("glob" in (c.reason or "").lower() for c in run.checks)


def test_a_sheet_that_is_not_in_the_workbook_fails():
    run = DataContract(data_contract_str=contract(WORKBOOK, physicalName="Shipments")).test()

    assert run.result != "passed"
    assert any("Shipments" in (c.reason or "") for c in run.checks)


def test_a_value_of_the_wrong_type_fails_only_its_column(tmp_path: Path):
    run = DataContract(data_contract_str=contract(edited(tmp_path, E3="n/a"), models=["orders"])).test()

    assert failed(run) == {"Check that field amount has type number": "Actual custom_sql(amount) was 1, expected = 0"}


def test_rows_below_an_empty_row_are_checked(tmp_path: Path):
    workbook = load_workbook(WORKBOOK)
    workbook["Orders"].insert_rows(3)
    workbook["Orders"]["E5"] = "n/a"
    workbook.save(tmp_path / "delivery.xlsx")

    run = DataContract(data_contract_str=contract(str(tmp_path / "delivery.xlsx"), models=["orders"])).test()

    # the empty row itself is skipped, or its order_id and order_date would be missing
    assert failed(run) == {"Check that field amount has type number": "Actual custom_sql(amount) was 1, expected = 0"}


def titled(tmp_path: Path, **cells) -> str:
    """The fixture workbook with a title row above the header of its Orders sheet, e.g. E4="n/a"."""
    workbook = load_workbook(WORKBOOK)
    orders = workbook["Orders"]
    orders.insert_rows(1)
    orders["A1"] = "Orders export"
    for cell, value in cells.items():
        orders[cell] = value
    workbook.save(tmp_path / "delivery.xlsx")
    return str(tmp_path / "delivery.xlsx")


XLSX_RANGE = [{"property": "xlsxRange", "value": "A2:H"}]


def test_a_sheet_with_its_header_below_a_title_is_read_from_its_range(tmp_path: Path):
    contract_str = contract(titled(tmp_path), models=["orders"], customProperties=XLSX_RANGE)

    assert DataContract(data_contract_str=contract_str).test().result == "passed"
    assert rows(contract_str, "orders") == rows(contract(WORKBOOK, models=["orders"]), "orders")


def test_a_value_of_the_wrong_type_is_found_within_the_range(tmp_path: Path):
    contract_str = contract(titled(tmp_path, E4="n/a"), models=["orders"], customProperties=XLSX_RANGE)

    assert failed(DataContract(data_contract_str=contract_str).test()) == {
        "Check that field amount has type number": "Actual custom_sql(amount) was 1, expected = 0"
    }


def test_a_sheet_with_its_header_below_a_title_fails_without_a_range(tmp_path: Path):
    run = DataContract(data_contract_str=contract(titled(tmp_path), models=["orders"])).test()

    assert "Check that field 'order_id' is present" in failed(run)


def test_a_sheet_with_an_empty_row_above_its_header_fails(tmp_path: Path):
    workbook = load_workbook(WORKBOOK)
    workbook["Orders"].insert_rows(1)
    workbook.save(tmp_path / "delivery.xlsx")

    run = DataContract(data_contract_str=contract(str(tmp_path / "delivery.xlsx"), models=["orders"])).test()

    assert run.result != "passed"
    assert any("row 1 of sheet 'Orders'" in (c.reason or "") for c in run.checks)


def test_a_range_that_starts_on_an_empty_row_fails(tmp_path: Path):
    path = titled(tmp_path)
    workbook = load_workbook(path)
    workbook["Orders"].insert_rows(2)
    workbook.save(path)

    run = DataContract(data_contract_str=contract(path, models=["orders"], customProperties=XLSX_RANGE)).test()

    assert run.result != "passed"
    assert any("row 2 of sheet 'Orders'" in (c.reason or "") for c in run.checks)


def test_a_fraction_in_an_integer_column_is_not_rounded(tmp_path: Path):
    run = DataContract(data_contract_str=contract(edited(tmp_path, F3=1.5), models=["orders"])).test()

    assert failed(run) == {
        "Check that field quantity has type integer": "Actual custom_sql(quantity) was 1, expected = 0"
    }


def test_a_header_is_matched_regardless_of_case(tmp_path: Path):
    contract_str = contract(edited(tmp_path, A1="ORDER_ID"), models=["orders"])

    assert [row[0] for row in rows(contract_str, "orders")] == ["A1", "A2", "A3"]


def test_a_date_that_is_not_a_date_fails_its_column(tmp_path: Path):
    run = DataContract(data_contract_str=contract(edited(tmp_path, C3="soon"), models=["orders"])).test()

    # as for JSON, the value reads as NULL, so the column being required reports it too
    assert failed(run) == {
        "Check that field order_date has type date": "Actual custom_sql(order_date) was 1, expected = 0",
        "Check that field order_date has no missing values": "Actual missing_count(order_date) was 1, expected = 0",
    }


def test_dates_below_a_first_cell_that_is_not_a_date_are_still_dates(tmp_path: Path):
    # DuckDB infers the type of a column from its first row, so these dates arrive as Excel serial numbers
    contract_str = contract(edited(tmp_path, D2="unknown"), models=["orders"])

    assert [row[3] for row in rows(contract_str, "orders")] == [
        None,
        dt.datetime(2026, 1, 3, 8, 0),
        dt.datetime(2026, 1, 5, 0, 0),
    ]


def test_a_date_in_a_text_column_reads_as_the_date_not_its_serial_number(tmp_path: Path):
    contract_str = contract(edited(tmp_path, H2=dt.date(2026, 2, 1)), models=["orders"])

    assert [row[7] for row in rows(contract_str, "orders")] == ["2026-02-01", None, "third"]


def test_a_column_that_is_not_in_the_sheet_fails_its_presence_check(tmp_path: Path):
    run = DataContract(data_contract_str=contract(edited(tmp_path, F1="qty"), models=["orders"])).test()

    assert "Check that field 'quantity' is present" in failed(run)


def test_a_nested_property_gets_no_type_check():
    data = yaml.safe_load(contract(WORKBOOK, models=["orders"]))
    data["schema"][0]["properties"].append(
        {"name": "meta", "logicalType": "object", "properties": [{"name": "source", "logicalType": "integer"}]}
    )

    run = DataContract(data_contract_str=yaml.safe_dump(data)).test()

    # a sheet has no nested columns, so such a check could only ever pass
    assert not any("meta.source" in check.name and "has type" in check.name for check in run.checks)


@pytest.mark.parametrize("sheet", ["Bob's Orders", 'Orders"; DROP TABLE x; --'])
def test_a_sheet_name_is_not_run_as_sql(tmp_path: Path, sheet):
    run = DataContract(data_contract_str=contract(edited(tmp_path, sheet_title=sheet), physicalName=sheet)).test()

    assert run.result == "passed", failed(run)


def test_an_untrusted_contract_reads_its_workbook():
    # The sandbox of an untrusted contract stops DuckDB from loading extensions on demand
    run = DataContract(data_contract_file=CONTRACT, untrusted_contract=True, inline_references=False).test()

    assert run.result == "passed", failed(run)
